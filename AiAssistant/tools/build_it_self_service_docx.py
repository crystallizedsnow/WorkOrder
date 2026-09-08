from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "src/main/resources/knowledge-base/计算机运维自助排障知识库.docx"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
PALE_GREEN = "EAF4EA"
PALE_RED = "FCE8E6"
GRAY = "666666"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_field(paragraph, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=GRAY)


def add_header_footer(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = "工单系统知识库  |  计算机运维自助排障"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_run_font(run, size=9, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("内部使用  ·  第 ")
    set_run_font(r, size=9, color=GRAY)
    add_field(footer, "PAGE")
    r = footer.add_run(" 页")
    set_run_font(r, size=9, color=GRAY)


def add_callout(doc, label, text, fill=LIGHT_BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), DARK_BLUE if fill != PALE_RED else "C5221F")
    borders.append(left)
    p_pr.append(borders)
    lead = p.add_run(label + "：")
    set_run_font(lead, bold=True, color=DARK_BLUE)
    body = p.add_run(text)
    set_run_font(body)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_topic(doc, title, symptoms, steps, checks, escalate, ticket_info, keywords):
    doc.add_heading(title, level=1)
    doc.add_heading("适用现象", level=2)
    add_bullets(doc, symptoms)
    doc.add_heading("自助处理步骤", level=2)
    add_steps(doc, steps)
    doc.add_heading("如何判断已经解决", level=2)
    add_bullets(doc, checks)
    doc.add_heading("停止自助并创建工单", level=2)
    add_callout(doc, "立即升级", "；".join(escalate) + "。", PALE_RED)
    doc.add_heading("创建工单时请提供", level=2)
    add_bullets(doc, ticket_info)
    p = doc.add_paragraph()
    r = p.add_run("检索关键词：")
    set_run_font(r, bold=True, color=GRAY)
    set_run_font(p.add_run("、".join(keywords)), color=GRAY)


def build():
    doc = Document()
    style_document(doc)
    add_header_footer(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("计算机运维自助排障知识库")
    set_run_font(r, size=26, bold=True, color=DARK_BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("员工自助指南 · 能安全解决就不创建工单，不能安全解决就及时升级")
    set_run_font(r, size=13, color=GRAY)

    meta = doc.add_table(rows=5, cols=2)
    meta.rows[0].cells[0].text = "属性"
    meta.rows[0].cells[1].text = "内容"
    tr_pr = meta.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for cell in meta.rows[0].cells:
        set_cell_shading(cell, BLUE)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, bold=True, color=WHITE)
    values = [
        ("知识版本", "2026.09.04.1"),
        ("适用范围", "公司管理的 Windows 电脑及常用办公服务"),
        ("目标用户", "普通员工，不要求管理员权限"),
        ("维护建议", "由 IT 运维团队审核后发布到生产知识库"),
    ]
    for row, (label, value) in zip(meta.rows[1:], values):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for run in row.cells[0].paragraphs[0].runs:
            set_run_font(run, bold=True, color=DARK_BLUE)
        for run in row.cells[1].paragraphs[0].runs:
            set_run_font(run)
    set_table_geometry(meta, [2700, 6660])

    doc.add_heading("先判断：现在是否适合自助处理", level=1)
    add_callout(doc, "自助原则", "只执行本文列出的低风险操作。不要拆机、修改注册表、关闭安全软件、绕过访问控制、安装来源不明的软件，或删除不确定的数据。", PALE_GREEN)
    doc.add_heading("可以先自助处理", level=2)
    add_bullets(doc, [
        "只有你自己的设备受影响，且没有冒烟、异响、进液、明显过热等硬件风险。",
        "问题属于网络重连、应用重启、缓存刷新、外设检查、磁盘整理等可逆操作。",
        "你仍能正常登录设备，重要文件已有保存，操作不会影响其他用户或生产系统。",
    ])
    doc.add_heading("以下情况直接创建工单或联系安全人员", level=2)
    add_bullets(doc, [
        "疑似账号被盗、收到异常验证码、误点钓鱼链接、发现勒索提示或安全软件报警。",
        "重要文件丢失、磁盘无法识别、电脑无法启动、反复蓝屏、冒烟、焦味、进液或电池鼓包。",
        "多人同时无法访问同一系统、网络或共享服务，可能是范围性故障。",
        "需要管理员权限、软件授权、访问权限、白名单、数据恢复或设备维修。",
    ])

    doc.add_heading("通用的三分钟检查", level=1)
    add_steps(doc, [
        "保存正在编辑的文件，记录报错原文和发生时间；不要只截取报错的一小部分。",
        "确认问题范围：只有一个应用、整台电脑，还是周围同事也受影响。",
        "关闭并重新打开故障应用；仍无效时，在确认文件已保存后正常重启电脑。",
        "重试一次并记录结果。不要连续重复登录、反复提交或频繁强制关机。",
    ])

    add_topic(doc, "无法访问互联网或公司网页", [
        "浏览器提示无网络、无法解析地址或连接超时。", "其他应用也可能无法联网，但电脑能够正常登录。"
    ], [
        "查看任务栏网络图标，确认没有启用飞行模式。", "使用 Wi-Fi 时断开当前网络，等待10秒后重新连接正确的公司网络；使用网线时重新插紧电脑端和墙面端。",
        "打开一个已知可用的网站确认是否为单个网站故障。", "在文件已保存的前提下正常重启电脑和浏览器。"
    ], ["至少两个不同网站可以正常打开。", "公司网页能够加载，且刷新后不再反复断线。"], [
        "周围多人同时断网", "网络端口或设备有烧焦、进液等异常", "反复断线超过15分钟", "需要修改代理、DNS或网卡驱动"
    ], ["有线或Wi-Fi连接方式及网络名称", "受影响网站地址和完整报错", "问题开始时间及同事是否受影响", "电脑资产编号"], ["无法上网", "网络断开", "网页打不开", "DNS", "连接超时"])

    add_topic(doc, "VPN 无法连接公司内网", [
        "外网正常，但 VPN 显示连接失败、认证失败或连接后仍打不开内网。"
    ], [
        "先确认普通互联网可用，并校准系统日期、时间和时区。", "完全退出 VPN 客户端后重新打开，确认选择公司规定的连接入口。",
        "重新输入自己的账号并完成合法的多因素认证；不要让他人代收验证码。", "连接成功后等待30秒，再访问一个已知内网地址。"
    ], ["VPN 状态显示已连接。", "已知内网页面可以打开，业务应用能够正常认证。"], [
        "收到本人未发起的验证码或登录提醒", "账号被锁定", "证书过期或客户端要求管理员权限", "多名同事同时连接失败"
    ], ["VPN客户端名称和版本", "错误代码及截图", "所在网络环境和发生时间", "外网是否正常、同事是否受影响"], ["VPN连接失败", "内网打不开", "认证失败", "证书过期", "多因素认证"])

    add_topic(doc, "账号密码或多因素认证异常", [
        "忘记密码、密码过期、登录失败，或正常登录时收不到验证码。"
    ], [
        "确认键盘大小写、输入法和用户名格式正确。", "只使用公司官方自助密码重置入口；通过收藏夹或公司门户进入，不点击陌生邮件中的重置链接。",
        "检查手机时间、网络和通知权限，然后仅重新请求一次验证码。", "重置后等待系统提示的同步时间，再尝试登录一次。"
    ], ["新密码可以登录，且没有异常登录提醒。", "多因素认证能够由本人正常完成。"], [
        "收到未主动发起的验证码", "怀疑密码泄露或误点钓鱼链接", "账号锁定或自助入口无法验证身份", "需要开通、变更或提升权限"
    ], ["账号标识，禁止填写密码和验证码", "登录的系统名称", "完整错误代码和发生时间", "是否收到异常登录或验证码"], ["忘记密码", "账号锁定", "收不到验证码", "MFA", "权限申请"])

    add_topic(doc, "电脑运行缓慢或应用无响应", [
        "开机慢、切换窗口卡顿、应用显示未响应，但电脑仍可操作。"
    ], [
        "保存能够保存的文件，关闭不再使用的浏览器标签页和大型应用。", "打开任务管理器观察 CPU、内存和磁盘是否持续接近100%，只记录占用最高的应用，不结束不认识的系统进程。",
        "确认系统盘至少保留约10%的可用空间。", "正常重启电脑；重启后只打开必要应用并验证。"
    ], ["重启后常用应用可以在合理时间内打开。", "CPU、内存或磁盘占用不再长时间接近100%。"], [
        "反复蓝屏或自动重启", "机身明显过热、异响或焦味", "不认识的程序持续高占用", "问题连续出现并影响工作"
    ], ["电脑资产编号和系统版本", "发生时间、频率及正在运行的应用", "任务管理器截图", "是否出现蓝屏、异响或过热"], ["电脑卡顿", "运行缓慢", "应用未响应", "CPU过高", "内存不足"])

    add_topic(doc, "浏览器或办公应用异常", [
        "单个网页显示异常、应用打不开、频繁闪退，或功能按钮没有响应。"
    ], [
        "保存文件并完全退出应用，再重新打开。", "确认只有一个网站或文件异常；换一个已知正常网页或空白文档验证。",
        "浏览器问题可使用无痕窗口测试；若无痕窗口正常，清理该网站的缓存和Cookie后重试。", "通过公司软件中心检查是否有批准的更新，不从搜索结果下载未知安装包。"
    ], ["应用重新打开后不再闪退。", "同一功能在正常窗口中可以稳定使用。"], [
        "文件包含重要数据且提示损坏", "应用要求管理员权限或许可证", "安全软件阻止运行", "多人同时遇到同一业务系统故障"
    ], ["应用名称和版本", "文件类型或网站地址", "报错截图和复现步骤", "换文件、换浏览器或无痕窗口的测试结果"], ["浏览器异常", "Office打不开", "应用闪退", "网页显示错误", "缓存Cookie"])

    add_topic(doc, "邮件或即时通讯收发异常", [
        "邮件长期停留在发件箱、收不到新邮件，或即时通讯消息无法发送。"
    ], [
        "确认网络正常，并查看服务是否显示离线或正在连接。", "检查发件箱中是否有超大附件；将大文件改用公司批准的文件分享方式。",
        "刷新或重新启动客户端，再通过网页版验证是否为本机客户端问题。", "检查垃圾邮件、会话筛选和搜索条件，但不要关闭公司安全过滤。"
    ], ["测试消息可以发送并由接收方确认收到。", "网页版和客户端状态一致。"], [
        "疑似钓鱼、冒充领导或异常附件", "多人同时无法收发", "邮箱容量异常且无法自行清理", "邮件包含敏感信息并被误发"
    ], ["邮件或通讯工具名称", "发送方、接收方和时间，不粘贴敏感正文", "错误提示和附件大小", "网页版测试结果"], ["邮件发不出", "收不到邮件", "发件箱", "消息发送失败", "钓鱼邮件"])

    add_topic(doc, "打印机无法打印", [
        "打印任务没有输出、队列暂停、打印机显示脱机或卡纸。"
    ], [
        "确认选择了正确的办公区域打印机，纸张和耗材提示正常。", "查看设备面板；若有明确卡纸指引，只按图示轻缓取出可见纸张，不拆卸盖板之外的部件。",
        "取消自己重复提交的打印任务，等待30秒后只发送一页测试。", "关闭并重新打开正在打印的应用；共享打印机不要擅自断电重启。"
    ], ["一页测试能够正常输出。", "打印队列不再持续累积，设备不再显示脱机。"], [
        "纸张撕裂且无法完整取出", "设备冒烟、异响、漏粉或有焦味", "多人均无法使用同一打印机", "需要安装驱动或管理员权限"
    ], ["打印机名称、位置和设备编号", "面板错误代码和照片", "文件类型及是否能打印测试页", "是否只有本人受影响"], ["打印失败", "打印机脱机", "打印队列", "卡纸", "无法连接打印机"])

    add_topic(doc, "磁盘空间不足", [
        "系统提示磁盘空间不足，应用无法更新或文件无法保存。"
    ], [
        "先清空回收站，并删除确认不再需要的下载文件和临时导出文件。", "将个人工作文件移动到公司批准的云盘或共享位置，确认同步成功后再删除本地副本。",
        "使用系统自带的存储清理功能处理临时文件；不要使用来源不明的清理软件。", "不要删除 Windows、Program Files、系统恢复或不认识的业务目录。"
    ], ["系统盘可用空间恢复到约10%以上。", "应用能够正常保存和更新。"], [
        "不知道大文件是否可以删除", "磁盘容量突然快速减少", "磁盘报错、无法识别或发出异响", "涉及业务数据迁移或恢复"
    ], ["电脑资产编号", "各磁盘总容量和剩余空间截图", "近期是否下载或生成大文件", "是否出现磁盘错误或异响"], ["C盘满", "磁盘空间不足", "无法保存", "清理临时文件", "存储空间"])

    add_topic(doc, "摄像头、麦克风或耳机不可用", [
        "会议中没有声音、对方听不到、摄像头黑屏，或耳机没有被识别。"
    ], [
        "确认耳机、摄像头已插紧，设备上的物理静音键或摄像头遮挡已关闭。", "在会议应用设置中选择正确的扬声器、麦克风和摄像头，不要只看系统默认设备。",
        "关闭其他可能占用摄像头或麦克风的应用，再重新加入会议。", "使用会议应用的测试功能录制几秒测试音频，不录制敏感会议内容。"
    ], ["测试音频可以播放，音量正常。", "本地预览画面正常，重新入会后对方可以接收。"], [
        "设备物理损坏或接口松动", "系统完全无法识别设备", "需要安装驱动或管理员权限", "多间会议室设备同时异常"
    ], ["设备品牌和连接方式", "会议应用名称和版本", "系统能否识别该设备", "错误截图及已测试的输入输出设备"], ["麦克风没声音", "摄像头黑屏", "耳机没声音", "会议设备", "设备未识别"])

    doc.add_heading("创建高质量工单", level=1)
    add_callout(doc, "不要提交敏感信息", "工单中不得填写密码、短信验证码、私钥、访问令牌或完整个人敏感数据。", PALE_RED)
    doc.add_heading("建议标题", level=2)
    p = doc.add_paragraph("[设备或系统] + [现象] + [影响范围]，例如：VPN认证失败，仅本人受影响")
    doc.add_heading("必须信息", level=2)
    add_bullets(doc, [
        "发生时间、是否持续、能否稳定复现。", "设备资产编号、操作系统、应用名称和版本。", "完整错误代码或清晰截图。",
        "影响范围：仅本人、同一地点多人，还是全公司。", "已经执行的自助步骤及每一步结果。", "业务影响和期望恢复时间。"
    ])

    doc.add_heading("知识维护说明", level=1)
    add_bullets(doc, [
        "本文只提供低风险、可逆的普通用户操作，不替代公司安全、权限和数据管理制度。",
        "公司专用系统地址、密码策略、VPN入口和SLA应由IT团队以权威来源补充，不在通用文档中猜测。",
        "每次更新后应执行意图路由、检索召回、旧版本泄漏、无答案拒答和回答忠实度评测。",
        "当本文与实时服务状态或IT人员指令冲突时，以实时状态和正式通知为准。",
    ])

    core = doc.core_properties
    core.title = "计算机运维自助排障知识库"
    core.subject = "工单系统员工自助排障知识"
    core.author = "工单系统 IT 运维知识库"
    core.keywords = "计算机运维, 自助排障, 工单, 网络, VPN, 账号, Office, 打印机"
    core.comments = "候选知识文档；须经IT审核、评测和发布门禁后进入生产知识库。"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
